"""
Tests for the deep constraint-analysis layer: diagnose_constraints,
repair_constraints, the Excel/Word report byte-builders, multi-turn analysis
memory (run_analysis save_as/recall + list_analysis_memory), and the chat_ui
executor wiring for all of the above.

Pure-function + fake-session tests only — no Bedrock, no network. Reuses the
shared streamlit stub installed by tests/conftest.py.
"""
import io
import json

import pandas as pd
import pytest

from components.chatbot import tools as T


# ==================== fixtures / helpers ====================

@pytest.fixture
def diag_df():
    """Data shaped to trigger each diagnosis: an over-subscribed scope, a tiny
    pool, and a healthy scope. NYC/CD = 100 containers; SEA/TL = 1 (ATMI only);
    LAX/CD = 200 (balanced)."""
    rows = []

    def add(carrier, port, cat, lane, n):
        rows.append({
            "Week Number": 10, "Category": cat, "Discharged Port": port,
            "Dray SCAC(FL)": carrier, "Facility": "FAC1", "Terminal": "T1",
            "SSL": "CMDU", "Vessel": "V1", "Lane": lane,
            "Container Count": n, "Base Rate": 100.0, "Total Rate": 100.0 * n,
            "Performance_Score": 0.9,
        })

    add("HJBT", "NYC", "CD", "USNYCABE8", 60)
    add("RKNE", "NYC", "CD", "USNYCABE8", 25)
    add("FRQT", "NYC", "CD", "USNYCABE8", 15)
    add("ATMI", "SEA", "TL", "USSEAXXX1", 1)
    add("ATMI", "LAX", "CD", "USLAXAAA1", 120)
    add("FRQT", "LAX", "CD", "USLAXAAA1", 80)
    return pd.DataFrame(rows)


def _c(**kw):
    """Build a full constraint row (spaced columns) from shorthand kwargs."""
    row = {col: None for col in T.CONSTRAINT_COLUMNS}
    alias = {"prio": "Priority Score", "carrier": "Carrier", "cat": "Category",
             "lane": "Lane", "port": "Port", "week": "Week Number",
             "max": "Maximum Container Count", "min": "Minimum Container Count",
             "pct": "Percent Allocation", "exfc": "Excluded FC",
             "terminal": "Terminal", "ssl": "SSL", "vessel": "Vessel"}
    for k, v in kw.items():
        row[alias.get(k, k)] = v
    row["_origin"] = "uploaded"
    return row


class _FakeSS(dict):
    def __getattr__(self, k):
        return self.get(k)

    def __setattr__(self, k, v):
        self[k] = v

    def setdefault(self, k, d=None):
        return super().setdefault(k, d)


def _run_with_fake_state(df, ss_init, name, tool_input):
    import components.chatbot.chat_ui as ui
    ss = _FakeSS(ss_init)
    orig = ui.st.session_state
    ui.st.session_state = ss
    try:
        execute = ui._make_tool_executor(df, None, "Base Rate")
        result, is_error = execute(name, tool_input)
        return result, is_error, ss
    finally:
        ui.st.session_state = orig


# ==================== diagnose_constraints ====================

def test_diagnose_empty_set_not_analyzable():
    out = T.diagnose_constraints([], None)
    assert out["analyzable"] is False


def test_diagnose_flags_oversubscription(diag_df):
    # NYC/CD pool=100; HJBT max 60 + RKNE 80% + FRQT 20% requests 160 > 100.
    cons = [
        _c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
        _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80),
        _c(prio=9, carrier="FRQT", port="NYC", cat="CD", pct=20),
    ]
    diag = T.diagnose_constraints(cons, diag_df)
    over = diag["over_subscribed_scopes"]
    assert any(s["port"] == "NYC" and s["category"] == "CD" for s in over)
    nyc = next(s for s in over if s["port"] == "NYC")
    assert nyc["available_containers"] == 100
    assert nyc["total_requested_containers"] > 100


def test_diagnose_flags_tiny_pool(diag_df):
    cons = [
        _c(prio=10, carrier="ATMI", port="SEA", cat="TL", pct=50),
        _c(prio=9, carrier="FRQT", port="SEA", cat="TL", pct=50),
        _c(prio=8, carrier="HDDR", port="SEA", cat="TL", min=5),
        _c(prio=8, carrier="RKNE", port="SEA", cat="TL", min=5),
    ]
    diag = T.diagnose_constraints(cons, diag_df)
    sea = [t for t in diag["tiny_pools"] if t["port"] == "SEA" and t["category"] == "TL"]
    assert sea, "SEA/TL should be flagged as a tiny pool"
    assert sea[0]["available_containers"] == 1
    assert sea[0]["carriers_present"] == ["ATMI"]


def test_diagnose_splits_dead_fixable_vs_acceptable(diag_df):
    cons = [
        _c(prio=10, carrier="ATMI", port="NYC", cat="CD", lane="ZZZ9", max=5),  # fixable typo
        _c(prio=9, carrier="FRQT", port="ORF", cat="CD", pct=50),               # acceptable out-of-scope
    ]
    diag = T.diagnose_constraints(cons, diag_df)
    fixable = diag["dead_scopes"]["fixable"]
    acceptable = diag["dead_scopes"]["acceptable"]
    assert any(d.get("Lane") == "ZZZ9" for d in fixable)
    assert any(d.get("Port") == "ORF" for d in acceptable)


def test_diagnose_lockout_not_counted_as_demand(diag_df):
    cons = [
        _c(prio=10, carrier="ATMI", port="LAX", cat="CD", pct=50),
        _c(prio=8, carrier="HJBT", port="LAX", cat="CD", pct=0),  # lockout
    ]
    diag = T.diagnose_constraints(cons, diag_df)
    lax = next(s for s in diag["scopes"] if s["port"] == "LAX")
    assert lax["percent_requested"] == 50


def test_diagnose_is_json_serializable(diag_df):
    cons = [_c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
            _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80)]
    json.dumps(T.diagnose_constraints(cons, diag_df))  # must not raise


# ==================== repair_constraints ====================

def test_repair_resolves_oversubscription_and_roundtrips(diag_df):
    from components.constraints.processor import (
        process_constraints_file, apply_constraints_to_data,
    )
    cons = [
        _c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
        _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80),
        _c(prio=9, carrier="FRQT", port="NYC", cat="CD", pct=20),
    ]
    rep = T.repair_constraints(cons, diag_df)
    assert rep["summary"]["rescaled"] >= 1
    for r in rep["constraints"]:
        assert not r.get("_problems")
        pct = r.get("Percent Allocation")
        if pct is not None:
            assert 0 <= pct <= 100
    clean = [r for r in rep["constraints"] if not r.get("_problems")]
    buf = io.BytesIO(T.constraints_to_excel_bytes(clean))
    reparsed = process_constraints_file(buf)
    apply_constraints_to_data(diag_df, reparsed, None)  # must not raise
    diag2 = T.diagnose_constraints(T.constraints_from_dataframe(reparsed), diag_df)
    assert not any(s["port"] == "NYC" and s["category"] == "CD"
                   for s in diag2["over_subscribed_scopes"])


def test_repair_preserves_lockouts_and_drops_dead(diag_df):
    cons = [
        _c(prio=10, carrier="HJBT", port="LAX", cat="CD", pct=0),               # lockout: keep
        _c(prio=9, carrier="ATMI", port="NYC", cat="CD", lane="ZZZ9", max=5),   # dead: drop
        _c(prio=8, carrier="FRQT", port="LAX", cat="CD", pct=50),               # healthy: keep
    ]
    rep = T.repair_constraints(cons, diag_df)
    kept = rep["constraints"]
    assert any(c.get("Carrier") == "HJBT" and c.get("Percent Allocation") == 0 for c in kept)
    assert not any(c.get("Lane") == "ZZZ9" for c in kept)
    assert rep["summary"]["dropped"] >= 1


def test_repair_collapses_tiny_pool_to_present_carriers(diag_df):
    cons = [
        _c(prio=10, carrier="ATMI", port="SEA", cat="TL", pct=100),
        _c(prio=9, carrier="HDDR", port="SEA", cat="TL", min=5),
        _c(prio=9, carrier="RKNE", port="SEA", cat="TL", min=5),
    ]
    rep = T.repair_constraints(cons, diag_df)
    kept_carriers = {c.get("Carrier") for c in rep["constraints"]}
    assert "ATMI" in kept_carriers
    assert "HDDR" not in kept_carriers and "RKNE" not in kept_carriers
    assert rep["summary"]["collapsed"] >= 2


def test_repair_empty_set_safe():
    out = T.repair_constraints([], None)
    assert out["constraints"] == []
    assert out["summary"]["rescaled"] == 0
    # Empty-set summary must carry the same keys as the normal path (schema parity).
    assert {"kept", "original_count", "clamped"} <= set(out["summary"])


def test_repair_clamps_out_of_range_amounts(diag_df):
    # A negative or >100 percent (which the rescale pass skips) must be CLAMPED to a
    # valid value by repair — handing back a flagged-but-invalid row is not "fixed".
    cons = [
        _c(prio=10, carrier="ATMI", port="SEA", cat="CD", pct=130),
        _c(prio=9, carrier="FRQT", port="SEA", cat="CD", pct=-10),
        _c(prio=8, carrier="HJBT", port="SEA", cat="CD", max=-5),
    ]
    rep = T.repair_constraints(cons, diag_df)
    for r in rep["constraints"]:
        pct = r.get("Percent Allocation")
        if pct is not None:
            assert 0 <= pct <= 100, "percent left out of range after repair"
        mx = r.get("Maximum Container Count")
        if mx is not None:
            assert mx >= 0, "negative cap left after repair"
        assert not r.get("_problems"), "repaired row still has validation problems"
    assert rep["summary"]["clamped"] >= 2


def test_repair_keeps_lockout_on_dead_scope(diag_df):
    # A lockout (Max 0 / 0%) is explicit intent — keep it even when its scope value
    # is absent from this run's data (it binds once that volume loads).
    cons = [
        _c(prio=10, carrier="ATMI", port="SEA", lane="ZZZ9", max=0),   # lockout on dead lane
        _c(prio=9, carrier="FRQT", port="SEA", lane="ZZZ9", max=5),    # non-lockout dead: drop
    ]
    rep = T.repair_constraints(cons, diag_df)
    carriers = [c.get("Carrier") for c in rep["constraints"]]
    assert "ATMI" in carriers, "lockout on a dead scope was wrongly dropped"
    assert "FRQT" not in carriers, "non-lockout dead rule should have been dropped"


# ==================== report byte-builders ====================

def test_workbook_bytes_valid_multisheet(diag_df):
    cons = [_c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
            _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80)]
    diag = T.diagnose_constraints(cons, diag_df)
    rep = T.repair_constraints(cons, diag_df)
    xb = T.build_analysis_workbook_bytes(diag, constraints_after=rep["constraints"])
    assert isinstance(xb, bytes) and len(xb) > 0
    xl = pd.ExcelFile(io.BytesIO(xb))
    assert "Diagnosis Summary" in xl.sheet_names
    assert "Corrected Constraints" in xl.sheet_names


def test_workbook_neutralizes_formula_injection(diag_df):
    # Hostile constraint strings (=cmd, @SUM, +1, -1) must NOT become live Excel
    # formulas in the downloaded report (CSV/formula-injection class).
    import openpyxl
    cons = [_c(prio=10, carrier="=cmd|'/c calc'!A1", port="=DDE", cat="+1+1",
               lane="@SUM", pct=50),
            _c(prio=9, carrier="FRQT", port="LAX", cat="TL", pct=50)]
    cons = [T._normalize_constraint(c) for c in cons]
    diag = T.diagnose_constraints(cons, diag_df)
    rep = T.repair_constraints(cons, diag_df)
    xb = T.build_analysis_workbook_bytes(diag, constraints_after=rep["constraints"])
    wb = openpyxl.load_workbook(io.BytesIO(xb))
    formula_cells = [(ws.title, c.coordinate, c.value)
                     for ws in wb.worksheets for row in ws.iter_rows()
                     for c in row if c.data_type == "f"]
    assert not formula_cells, f"live formula cells leaked: {formula_cells}"


def test_report_builders_survive_control_chars(diag_df):
    # A C0 control char in a constraint value must not crash either builder.
    cons = [T._normalize_constraint(
        _c(prio=10, carrier="FRQT", port="LAX", cat="TL", lane="LA\x07X\x00bad", pct=50))]
    cons[0]["_origin"] = "uploaded"
    diag = T.diagnose_constraints(cons, diag_df)
    rep = T.repair_constraints(cons, diag_df)
    xb = T.build_analysis_workbook_bytes(diag, constraints_after=rep["constraints"])
    assert isinstance(xb, bytes) and len(xb) > 0
    pd.ExcelFile(io.BytesIO(xb))  # re-opens
    db = T.build_analysis_report_docx_bytes(diag, repair=rep)
    if db is not None:
        from docx import Document
        Document(io.BytesIO(db))  # parses


def test_docx_bytes_present_and_guarded(diag_df, monkeypatch):
    cons = [_c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
            _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80)]
    diag = T.diagnose_constraints(cons, diag_df)
    db = T.build_analysis_report_docx_bytes(diag)
    assert isinstance(db, bytes) and len(db) > 0
    import builtins
    real = builtins.__import__

    def fake(name, *a, **k):
        if name == "docx" or name.startswith("docx."):
            raise ImportError("simulated missing python-docx")
        return real(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake)
    assert T.build_analysis_report_docx_bytes(diag) is None


# ==================== analysis memory ====================

def test_list_analysis_memory_empty_and_populated():
    assert T.list_analysis_memory({})["count"] == 0
    mem = {"port_totals": pd.Series([1, 2, 3], name="x"),
           "top": pd.DataFrame({"a": [1]}), "n": 5}
    out = T.list_analysis_memory(mem)
    assert out["count"] == 3
    kinds = {r["name"]: r["kind"] for r in out["results"]}
    assert kinds == {"port_totals": "series", "top": "dataframe", "n": "scalar"}


def test_run_analysis_save_and_recall(diag_df):
    out1 = T.run_analysis(
        diag_df, "result = df.groupby('Discharged Port')['Container Count'].sum()",
        save_as="port_totals")
    assert out1["ok"] and out1["saved_as"] == "port_totals"
    assert out1["_save"]["name"] == "port_totals"
    mem = {out1["_save"]["name"]: out1["_save"]["value"]}
    out2 = T.run_analysis(diag_df, "result = memory['port_totals'].max()", memory=mem)
    assert out2["ok"]
    assert out2["result"]["value"] == 200  # LAX = 120 + 80


def test_run_analysis_memory_is_read_only(diag_df):
    s = pd.Series({"LAX": 200}, name="x")
    mem = {"t": s}
    T.run_analysis(diag_df, "memory['t']['LAX'] = -1\nresult = 1", memory=mem)
    assert int(s["LAX"]) == 200


def test_run_analysis_no_raw_leak_without_save_as(diag_df):
    out = T.run_analysis(diag_df, "result = df['Container Count'].sum()")
    assert "_save" not in out and "raw" not in out


# ==================== executor dispatch ====================

def test_executor_repair_stages_not_applies(diag_df):
    staged = [
        _c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
        _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80),
        _c(prio=9, carrier="FRQT", port="NYC", cat="CD", pct=20),
    ]
    result, is_error, ss = _run_with_fake_state(
        diag_df, {"chatbot_staged_constraints": staged,
                  "chatbot_applied_constraints": []},
        "repair_constraints", {})
    assert is_error is False
    assert len(ss["chatbot_staged_constraints"]) >= 1
    assert ss["chatbot_applied_constraints"] == []
    assert "summary" in result


def test_executor_report_stashes_bytes_not_returns_them(diag_df):
    staged = [_c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
              _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80)]
    result, is_error, ss = _run_with_fake_state(
        diag_df, {"chatbot_staged_constraints": staged},
        "generate_analysis_report", {"include_fix": True})
    assert is_error is False
    assert result.get("report_ready") is True
    assert isinstance(ss["chatbot_report_xlsx_bytes"], bytes)
    assert not any(isinstance(v, bytes) for v in result.values())
    json.dumps(result)  # ack must be JSON-serializable


def test_executor_diagnose_routes(diag_df):
    staged = [_c(prio=10, carrier="HJBT", port="NYC", cat="CD", max=60),
              _c(prio=9, carrier="RKNE", port="NYC", cat="CD", pct=80)]
    result, is_error, _ = _run_with_fake_state(
        diag_df, {"chatbot_staged_constraints": staged},
        "diagnose_constraints", {})
    assert is_error is False
    assert result["analyzable"] is True


def test_executor_run_analysis_save_persists_in_session(diag_df):
    result, is_error, ss = _run_with_fake_state(
        diag_df, {"chatbot_analysis_memory": {}},
        "run_analysis",
        {"code": "result = df['Container Count'].sum()", "save_as": "total"})
    assert is_error is False
    assert "total" in ss["chatbot_analysis_memory"]
    out, _, _ = _run_with_fake_state(
        diag_df, {"chatbot_analysis_memory": ss["chatbot_analysis_memory"]},
        "list_analysis_memory", {})
    assert out["count"] == 1


def test_remember_analysis_evicts_oldest_past_cap():
    import components.chatbot.chat_ui as ui
    ss = _FakeSS({"chatbot_analysis_memory": {}})
    orig = ui.st.session_state
    ui.st.session_state = ss
    try:
        for i in range(ui._ANALYSIS_MEMORY_MAX + 3):
            ui._remember_analysis("r%d" % i, i)
        mem = ss["chatbot_analysis_memory"]
        assert len(mem) == ui._ANALYSIS_MEMORY_MAX
        assert "r0" not in mem
        assert ("r%d" % (ui._ANALYSIS_MEMORY_MAX + 2)) in mem
    finally:
        ui.st.session_state = orig
